from io import BytesIO
import pandas as pd
from openpyxl.styles import PatternFill
 
def export_excel(df, sheet_name="Data"):
 
    output = BytesIO()
 
    highlight_cols = [
        "CHR_WT",
        "TOTAL_FRT",
        "CPKG"
    ]
 
    fill = PatternFill(
        start_color="FFD6D6",
        end_color="FFD6D6",
        fill_type="solid"
    )
 
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
 
        df.to_excel(
            writer,
            index=False,
            sheet_name=sheet_name
        )
 
        ws = writer.sheets[sheet_name]
 
        col_idx = {
            col: idx + 1
            for idx, col in enumerate(df.columns)
        }
 
        for col in highlight_cols:
 
            if col in col_idx:
 
                c = col_idx[col]
 
                for row in range(2, len(df) + 2):
 
                    ws.cell(
                        row=row,
                        column=c
                    ).fill = fill
 
    output.seek(0)
 
    return output
 