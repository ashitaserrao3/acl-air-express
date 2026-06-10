import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl.styles import PatternFill, Font
from utils.dashboard import show_dashboard

def run():
    st.markdown("## 📦 EDS Processor")

    uploaded_files = st.file_uploader(
        "Upload EDS Excel Files",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        key="eds_upload"
    )

    def find_col(df, names):
        for col in df.columns:
            c = str(col).strip().upper()
            for n in names:
                if c == n.strip().upper():
                    return col
        return None

    def get_airline(f):
        if pd.isna(f):
            return None
        code = str(f).strip()[:2].upper()
        return {
            "AI": "AirIndia",
            "IX": "AirIndia",
            "6E": "Indigo",
            "SG": "SpiceJet",
            "QP": "Akasa"
        }.get(code, "Other")

    def get_slab(wt):
        try:
            wt = float(wt)
        except:
            return None
        if wt < 45: return "Less than 45Kg"
        elif wt < 100: return "45Kg +"
        elif wt < 250: return "100Kg +"
        elif wt < 300: return "250Kg +"
        elif wt < 500: return "300Kg +"
        elif wt < 1000: return "500Kg +"
        return "1000Kg +"

    if uploaded_files:

        final_df = pd.DataFrame()

        for file in uploaded_files:

            raw = pd.read_excel(file, header=None)

            headers = raw.iloc[2].astype(str).str.strip()
            df = raw.iloc[3:].copy()
            df.columns = headers
            df = df.reset_index(drop=True)

            awb_col = find_col(df, ["AWB NO.", "AWB NO"])
            if awb_col is None:
                continue

            df = df[df[awb_col].notna()]
            df = df[~df[awb_col].astype(str).str.contains("Airline", case=False, na=False)]

            date_col = find_col(df, ["DATE"])
            flight_col = find_col(df, ["FLIGHT NO.", "FLIGHT NO"])
            dest_col = find_col(df, ["DSTN"])
            pcs_col = find_col(df, ["PCS"])
            wt_col = find_col(df, ["CWGT KGS.", "CWGT KGS"])
            rate_col = find_col(df, ["RATE"])
            freight_col = find_col(df, ["FREIGHT CHS"])

            awb_do_col = find_col(df, ["AWB DO"])
            al_col = find_col(df, ["A/L CHG"])
            apt_out_col = find_col(df, ["APT OUT"])
            apt_in_col = find_col(df, ["APT IN"])
            svc_col = find_col(df, ["SVC CHG"])

            cgst_col = find_col(df, ["CGST 9%"])
            sgst_col = find_col(df, ["SGST 9%"])
            igst_col = find_col(df, ["IGST 18%"])
            amount_col = find_col(df, ["AMOUNT"])

            num_cols = [wt_col, rate_col, freight_col, awb_do_col, al_col,
                        apt_out_col, apt_in_col, svc_col, cgst_col,
                        sgst_col, igst_col, amount_col]

            for c in num_cols:
                if c is not None:
                    df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

            df["ADDON_CHR"] = (
                df[awb_do_col].fillna(0)
                + df[al_col].fillna(0)
                + df[apt_out_col].fillna(0)
                + df[apt_in_col].fillna(0)
                + df[svc_col].fillna(0)
            )

            df["TOTAL_FRT"] = df[freight_col] + df["ADDON_CHR"]

            df["CPKG"] = (df["TOTAL_FRT"] / df[wt_col]).round(2)
            df["ADDON_PER_KG"] = (df["ADDON_CHR"] / df[wt_col]).round(2)

            df["AWB_DATE"] = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True)

            parts = file.name.replace(".xlsx","").replace(".xls","").split("-")
            invoice_no = "-".join(parts[:3])
            origin = parts[0]

            out = pd.DataFrame()
            out["INVOICE_NO"] = invoice_no
            out["BILL_PERIOD"] = df["AWB_DATE"].dt.day.apply(lambda x: "FFN" if pd.notna(x) and x <= 15 else "SFN")
            out["AGENT"] = "EDS"
            out["TRNSPT_MODE"] = "AIR"
            out["ORIGIN"] = origin
            out["DEST"] = df[dest_col]
            out["OD_PAIR"] = out["ORIGIN"] + "-" + out["DEST"].astype(str)

            out["AWB_NO"] = df[awb_col]
            out["AWB_DATE"] = df["AWB_DATE"].dt.strftime("%d-%b-%Y")
            out["FLIGHT_NO"] = df[flight_col]
            out["AIRLINE"] = df[flight_col].apply(get_airline)

            out["PKGS"] = df[pcs_col]
            out["CHR_WT"] = df[wt_col]
            out["RATE"] = df[rate_col]
            out["BASIC_FRT"] = df[freight_col]

            out ["AWB DO"] = df[awb_do_col]
            out ["A/L CHG"] = df[al_col]
            out ["APT OUT"] = df[apt_out_col]
            out ["APT IN"] = df[apt_in_col]
            out ["SVC CHG"] = df[svc_col]

            out["TOTAL_FRT"] = df["TOTAL_FRT"]
            out["ADDON_CHR"] = df["ADDON_CHR"]
            out["ADDON_PER_KG"] = df["ADDON_PER_KG"]
            out["CPKG"] = df["CPKG"]
            out["CGST"] = df[cgst_col]
            out["SGST"] = df[sgst_col]
            out["IGST"] = df[igst_col]
            out["AMOUNT"] = df[amount_col]
            out["SLAB"] = df[wt_col].apply(get_slab)
            out["LODGE_MODE"] = df[awb_col].astype(str).apply(
                lambda x: "Direct" if "-" in x else "Console"
            )

            final_df = pd.concat([final_df, out], ignore_index=True)

        st.success(f"✅ {len(final_df)} rows processed")

        filtered_df = show_dashboard(final_df)

        st.dataframe(filtered_df, use_container_width=True)

        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            final_df.to_excel(writer, index=False, sheet_name="EDS")

            ws = writer.sheets["EDS"]

            header_fill = PatternFill(
                start_color="4F81BD",
                end_color="4F81BD",
                fill_type="solid"
            )

            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

        output.seek(0)

        st.download_button(
            "📥 Download Excel",
            data=output,
            file_name="EDS_INDEX_OUTPUT.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
