import streamlit as st
import pandas as pd
import openpyxl
import re

from io import BytesIO
from datetime import datetime
from openpyxl.styles import PatternFill

from utils.dashboard import show_dashboard


def run():
    st.markdown("## 📦FDC Processor")

    # ------------------ EXTRACT INVOICE ------------------ #
    def extract_invoice_details(df_raw):
        invoice_no = None
        invoice_date = None

        for _, row in df_raw.iterrows():
            text = " ".join(
                [str(x) for x in row if pd.notna(x)]
            ).lower()

            inv_match = re.search(
                r'([0-9]{1,3}/[0-9]{4}[-][0-9]{2})',
                text
            )

            if inv_match and not invoice_no:
                invoice_no = inv_match.group(1)

            dt_match = re.search(
                r'dt\s*[:-]?\s*(\d{2}[./-]\d{2}[./-]\d{4})',
                text
            )

            if dt_match and not invoice_date:
                invoice_date = dt_match.group(1)

        return invoice_no, invoice_date

    # ------------------ SHEET SELECTION ------------------ #
    def extract_dt_from_sheet(ws):
        try:
            for row in ws.iter_rows(
                min_row=1,
                max_row=10,
                values_only=True
            ):
                for cell in row:
                    if cell:
                        match = re.search(
                            r'(\d{2}.\d{2}.\d{4})',
                            str(cell)
                        )

                        if match:
                            return datetime.strptime(
                                match.group(1),
                                "%d.%m.%Y"
                            )

            return datetime(1900, 1, 1)

        except:
            return datetime(1900, 1, 1)

    def get_latest_sheet(wb):
        return max(
            wb.sheetnames,
            key=lambda s: extract_dt_from_sheet(wb[s])
        )

    # ------------------ HEADER ------------------ #
    def detect_header_row(df_raw):
        for idx, row in df_raw.iterrows():
            if any(
                "con" in str(cell).lower()
                for cell in row
            ):
                return idx

        return None

    # ------------------ COLUMN FINDER ------------------ #
    def find_col(df, keywords):
        for col in df.columns:
            col_clean = str(col).lower()

            for k in keywords:
                if k.lower() in col_clean:
                    return col

        return None

    # ------------------ SLAB ------------------ #
    def get_slab(w):
        if pd.isna(w):
            return None

        if w < 45:
            return "Less than 45Kg"
        elif w < 100:
            return "45Kg +"
        elif w < 250:
            return "100Kg +"
        elif w < 300:
            return "250Kg +"
        elif w < 500:
            return "300Kg +"
        elif w < 1000:
            return "500Kg +"

        return "1000Kg +"

    # ------------------ FILE UPLOAD ------------------ #
    uploaded_files = st.file_uploader(
        "Upload FDC Excel files",
        type=["xlsx"],
        accept_multiple_files=True,
        key="fdc_upload"
    )

    if uploaded_files:

        # =====================================================
        # CLEAR BUTTON
        # =====================================================
        #if st.button("Clear Processed Data"):
            #if "fdc_df" in st.session_state:
                #del st.session_state["fdc_df"]

            #st.rerun()

        # =====================================================
        # PROCESS ONLY ONCE
        # =====================================================
        if "fdc_df" not in st.session_state:

            combined_df = pd.DataFrame()
            progress = st.progress(0)

            for i, file in enumerate(uploaded_files):

                st.write(f"Processing file: {file.name}")

                wb = openpyxl.load_workbook(file)

                latest_sheet = get_latest_sheet(wb)
                sheet = wb[latest_sheet]

                st.write(f"Using sheet: {latest_sheet}")

                # ---- UNMERGE ---- #
                for m in list(sheet.merged_cells.ranges):
                    sheet.unmerge_cells(str(m))

                temp = BytesIO()
                wb.save(temp)
                temp.seek(0)

                df_raw = pd.read_excel(
                    temp,
                    sheet_name=latest_sheet,
                    header=None
                )

                invoice_no, invoice_date = extract_invoice_details(df_raw)

                header_row = detect_header_row(df_raw)

                if header_row is None:
                    st.warning(f"Header not found in {file.name}")
                    continue

                df = pd.read_excel(
                    temp,
                    sheet_name=latest_sheet,
                    header=header_row
                )

                # ---- CLEAN COLUMNS ---- #
                df.columns = (
                    df.columns
                    .astype(str)
                    .str.replace("\n", "")
                    .str.strip()
                    .str.upper()
                )

                # ---- COLUMN MAP ---- #
                awb_col = find_col(df, ["CON.NO", "AWB"])
                date_col = find_col(df, ["DATE"])
                dest_col = find_col(df, ["DEST"])
                wt_col = find_col(df, ["WEIGHT", "WT"])
                rate_col = find_col(df, ["RATE"])
                qty_col = find_col(df, ["QTY"])
                oda_col = find_col(df, ["ODA"])

                if not awb_col:
                    st.warning(f"AWB column missing in {file.name}")
                    continue

                # ---- BASE FIELDS ---- #
                df["AWB_NO"] = (
                    df[awb_col]
                    .astype(str)
                    .str.strip()
                )

                if date_col:
                    df["AWB_DATE"] = pd.to_datetime(
                        df[date_col]
                        .astype(str)
                        .str.strip(),
                        format="%d.%m.%Y",
                        errors="coerce"
                    )
                else:
                    df["AWB_DATE"] = None

                df["BILL_PERIOD"] = df["AWB_DATE"].apply(
                    lambda d: "FFN"
                    if pd.notna(d) and d.day <= 15
                    else "SFN"
                )

                df["DEST"] = (
                    df[dest_col]
                    if dest_col
                    else "UNKNOWN"
                )

                df["CHR_WT"] = pd.to_numeric(
                    df.get(wt_col),
                    errors="coerce"
                )

                df["RATE"] = pd.to_numeric(
                    df.get(rate_col),
                    errors="coerce"
                )

                df["PKGS"] = pd.to_numeric(
                    df.get(qty_col),
                    errors="coerce"
                )

                # ---- CALCULATIONS ---- #
                df["BASIC_FRT"] = (
                    df["RATE"] * df["CHR_WT"]
                )

                df["ODA_CHARGES"] = (
                    pd.to_numeric(
                        df.get(oda_col),
                        errors="coerce"
                    ).fillna(0)
                )

                df["TOTAL_FRT"] = (
                    df["BASIC_FRT"] +
                    df["ODA_CHARGES"]
                )

                df["ADDON_CHR"] = df["ODA_CHARGES"]

                df["CPKG"] = (
                    df["TOTAL_FRT"] /
                    df["CHR_WT"]
                ).round(2)

                df["ADDON_PER_KG"] = (
                    df["ADDON_CHR"] /
                    df["CHR_WT"]
                ).round(2)

                # ---- CLEAN ---- #
                df = df[
                    df["AWB_NO"].str.len() > 5
                ]

                # ---- STATIC ---- #
                df["AGENT"] = "FDC"
                df["TRNSPT_MODE"] = "ROAD"
                df["ORIGIN"] = "HYD"
                df["LODGE_MODE"] = "N/A"

                df["OD_PAIR"] = (
                    df["ORIGIN"] +
                    "-" +
                    df["DEST"].astype(str)
                )

                df["INVOICE_NO"] = (
                    invoice_no
                    if invoice_no
                    else file.name
                )

                df["AWB_DATE"] = (
                    df["AWB_DATE"]
                    .dt.strftime("%d-%b-%Y")
                )

                # ---- SLAB ---- #
                df["SLAB"] = (
                    df["CHR_WT"]
                    .apply(get_slab)
                )

                # ---- FINAL ---- #
                final_cols = [
                    "INVOICE_NO",
                    "BILL_PERIOD",
                    "AGENT",
                    "TRNSPT_MODE",
                    "OD_PAIR",
                    "ORIGIN",
                    "DEST",
                    "AWB_NO",
                    "AWB_DATE",
                    "PKGS",
                    "CHR_WT",
                    "RATE",
                    "BASIC_FRT",
                    "ODA_CHARGES",
                    "TOTAL_FRT",
                    "ADDON_CHR",
                    "ADDON_PER_KG",
                    "CPKG",
                    "SLAB",
                    "LODGE_MODE"
                ]

                for col in final_cols:
                    if col not in df.columns:
                        df[col] = None

                df = df[final_cols]

                combined_df = pd.concat(
                    [combined_df, df],
                    ignore_index=True
                )

                progress.progress(
                    (i + 1) / len(uploaded_files)
                )

            # ---- OUTPUT ---- #
            if combined_df.empty:
                st.warning("No valid data found.")
                return

            # STANDARD DASHBOARD COLUMNS
            if "TOTAL_FRT" not in combined_df.columns:
                combined_df["TOTAL_FRT"] = 0

            if "ORIGIN" not in combined_df.columns:
                combined_df["ORIGIN"] = "HYD"

            if "DEST" not in combined_df.columns:
                combined_df["DEST"] = "UNKNOWN"

            if "CHR_WT" not in combined_df.columns:
                combined_df["CHR_WT"] = 0

            # STORE SESSION
            st.session_state["fdc_df"] = combined_df

        # =====================================================
        # LOAD SESSION DATA
        # =====================================================
        combined_df = st.session_state["fdc_df"]

        st.success(
            f"✅ Combined {len(combined_df)} rows"
        )

        # =====================================================
        # DASHBOARD
        # =====================================================
        filtered_df = show_dashboard(combined_df)

        st.dataframe(
            filtered_df,
            use_container_width=True
        )

        # =====================================================
        # EXPORT
        # =====================================================
        output = BytesIO()

        fill = PatternFill(
            start_color="FFD6D6",
            end_color="FFD6D6",
            fill_type="solid"
        )

        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:

            combined_df.to_excel(
                writer,
                index=False,
                sheet_name="FDC"
            )

            ws = writer.sheets["FDC"]

            col_idx = {
                c: i + 1
                for i, c in enumerate(combined_df.columns)
            }

            highlight_cols = [
                "CHR_WT",
                "TOTAL_FRT",
                "CPKG"
            ]

            for col in highlight_cols:

                if col in col_idx:

                    col_pos = col_idx[col]

                    for row in range(
                        2,
                        len(combined_df) + 2
                    ):
                        ws.cell(
                            row=row,
                            column=col_pos
                        ).fill = fill

        output.seek(0)

        st.download_button(
            "📥 Download Combined",
            data=output,
            file_name="FDC_Combined.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )