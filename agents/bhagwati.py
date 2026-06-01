import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from utils.dashboard import show_dashboard


def run():

    st.title("📦BHAGWATI Processor")

    # ---------------- FUNCTIONS ---------------- #

    def unmerge_excel(file):

        file.seek(0)

        wb = load_workbook(
            file,
            data_only=True
        )

        for sheet in wb.worksheets:

            merged_ranges = list(
                sheet.merged_cells.ranges
            )

            for merged_range in merged_ranges:

                sheet.unmerge_cells(
                    str(merged_range)
                )

        output = BytesIO()

        wb.save(output)

        output.seek(0)

        return output

    def find_header_row(df):

        for i, row in df.iterrows():

            row_values = [
                str(x).strip().upper()
                for x in row.values
            ]

            if (
                "HAWB NO" in row_values
                and
                "CHG WT" in row_values
            ):
                return i

        return None

    def get_airline(flight):

        if pd.isna(flight):
            return None

        f = (
            str(flight)
            .strip()
            .upper()
        )

        if (
            "INDIGO" in f
            or f.startswith("6E")
        ):

            return "Indigo"

        elif (
            "AIRIN" in f
            or "AIR INDIA" in f
            or f.startswith("AI")
        ):

            return "AirIndia"

        elif (
            "AKASA" in f
            or f.startswith("QP")
        ):

            return "Akasa"

        else:
            return "Other"

    def get_slab(wt):

        try:

            wt = float(wt)

            if wt < 45:
                return "Less than 45Kg"

            elif wt < 100:
                return "45Kg +"

            elif wt < 250:
                return "100Kg +"

            elif wt < 500:
                return "250Kg +"

            elif wt < 1000:
                return "500Kg +"

            else:
                return "1000Kg +"

        except:
            return None

    def get_mode(hawb_charge):

        try:

            hawb_charge = float(hawb_charge)

            if round(hawb_charge, 0) == 50:
                return "Console"

            elif round(hawb_charge, 0) == 500:
                return "Direct"

            else:
                return "Other"

        except:
            return None

    # ---------------- FILE UPLOAD ---------------- #

    uploaded_files = st.file_uploader(
        "Upload Excel Files",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        key="bhagwati_upload"
    )

    if uploaded_files:

        # =====================================================
        # CLEAR BUTTON
        # =====================================================

        #if st.button("Clear Processed Data"):

            #if "bhagwati_df" in st.session_state:
                #del st.session_state["bhagwati_df"]

            #st.rerun()

        # =====================================================
        # PROCESS ONLY ONCE
        # =====================================================

        if "bhagwati_df" not in st.session_state:

            dfs = []

            for file in uploaded_files:

                st.write(
                    f"Processing: {file.name}"
                )

                cleaned_excel = unmerge_excel(file)

                raw_df = pd.read_excel(
                    cleaned_excel,
                    header=None
                )

                header_row = find_header_row(raw_df)

                if header_row is None:

                    st.warning(
                        f"Could not detect table header in {file.name}"
                    )

                    continue

                cleaned_excel.seek(0)

                df = pd.read_excel(
                    cleaned_excel,
                    header=header_row
                )

                df.columns = (
                    df.columns
                    .astype(str)
                    .str.strip()
                )

                if "HAWB NO" in df.columns:

                    df = df[
                        df["HAWB NO"].notna()
                    ]

                rename_map = {

                    "HAWB NO": "AWB_NO",
                    "DATE": "AWB_DATE",
                    "PEC.": "PKGS",
                    "CHG WT": "CHR_WT",
                    "DEST": "DEST",
                    "FLT NO.": "FLIGHT_NO",
                    "RATE": "RATE",
                    "FREIGHT": "BASIC_FRT",
                    "AMOUNT": "TOTAL_FRT",
                    "HAWB": "hawb_chr"
                }

                df = df.rename(
                    columns=rename_map
                )

                if "AWB_DATE" in df.columns:

                    df["AWB_DATE"] = pd.to_datetime(
                        df["AWB_DATE"],
                        dayfirst=True,
                        errors="coerce"
                    )

                numeric_cols = [
                    "CHR_WT",
                    "RATE",
                    "BASIC_FRT",
                    "hawb_chr",
                    "TOTAL_FRT"
                ]

                for col in numeric_cols:

                    if col in df.columns:

                        df[col] = pd.to_numeric(
                            df[col],
                            errors="coerce"
                        )

                if (
                    "TOTAL_FRT" not in df.columns
                    and
                    "BASIC_FRT" in df.columns
                    and
                    "hawb_chr" in df.columns
                ):

                    df["TOTAL_FRT"] = (
                        df["BASIC_FRT"]
                        +
                        df["hawb_chr"]
                    )

                if (
                    "TOTAL_FRT" in df.columns
                    and
                    "BASIC_FRT" in df.columns
                ):

                    df["ADDON_CHR"] = (
                        df["TOTAL_FRT"]
                        -
                        df["BASIC_FRT"]
                    )

                if (
                    "ADDON_CHR" in df.columns
                    and
                    "CHR_WT" in df.columns
                ):

                    df["ADDON_PER_KG"] = (
                        df["ADDON_CHR"]
                        /
                        df["CHR_WT"].replace(
                            0,
                            pd.NA
                        )
                    ).round(2)

                if (
                    "TOTAL_FRT" in df.columns
                    and
                    "CHR_WT" in df.columns
                ):

                    df["CPKG"] = (
                        df["TOTAL_FRT"]
                        /
                        df["CHR_WT"].replace(
                            0,
                            pd.NA
                        )
                    ).round(2)

                if "FLIGHT_NO" in df.columns:

                    df["AIRLINE"] = (
                        df["FLIGHT_NO"]
                        .apply(get_airline)
                    )

                if "CHR_WT" in df.columns:

                    df["SLAB"] = (
                        df["CHR_WT"]
                        .apply(get_slab)
                    )

                if "hawb_chr" in df.columns:

                    df["LODGE_MODE"] = (
                        df["hawb_chr"]
                        .apply(get_mode)
                    )

                filename = (
                    file.name
                    .replace(".xlsx", "")
                    .replace(".xls", "")
                )

                df["INVOICE_NO"] = filename

                df = df.dropna(how="all")

                dfs.append(df)

            # ---------------- FINAL OUTPUT ---------------- #

            if not dfs:

                st.warning(
                    "No valid data extracted"
                )

                return

            combined_df = pd.concat(
                dfs,
                ignore_index=True
            )

            combined_df = (
                combined_df
                .drop_duplicates()
            )

            # ---------------- FIXED COLUMNS ---------------- #

            combined_df["AGENT"] = "BHAGWATI"
            combined_df["TRNSPT_MODE"] = "Air"
            combined_df["ORIGIN"] = "PNQ"

            if "DEST" in combined_df.columns:

                combined_df["OD_PAIR"] = (

                    combined_df["ORIGIN"]
                    .astype(str)

                    +

                    "-"

                    +

                    combined_df["DEST"]
                    .astype(str)
                )

            if "AWB_DATE" in combined_df.columns:

                combined_df = combined_df[
                    combined_df["AWB_DATE"].notna()
                ]

                combined_df["BILL_PERIOD"] = (
                    combined_df["AWB_DATE"]
                    .apply(
                        lambda x:
                        "FFN"
                        if pd.notna(x)
                        and x.day <= 15
                        else "SFN"
                    )
                )

                combined_df["AWB_DATE"] = (
                    combined_df["AWB_DATE"]
                    .dt.strftime("%d-%b-%Y")
                )

            if "AWB_NO" in combined_df.columns:

                combined_df = combined_df[
                    combined_df["AWB_NO"]
                    .astype(str)
                    .str.upper()
                    != "TOTAL_FRT"
                ]

            # STANDARD DASHBOARD COLUMNS

            if "TOTAL_FRT" not in combined_df.columns:
                combined_df["TOTAL_FRT"] = 0

            if "ORIGIN" not in combined_df.columns:
                combined_df["ORIGIN"] = "PNQ"

            if "DEST" not in combined_df.columns:
                combined_df["DEST"] = "UNKNOWN"

            if "CHR_WT" not in combined_df.columns:
                combined_df["CHR_WT"] = 0

            # ---------------- COLUMN ORDER ---------------- #

            preferred_order1 = [

                "INVOICE_NO",
                "BILL_PERIOD",
                "AGENT",
                "TRNSPT_MODE",

                "OD_PAIR",
                "ORIGIN",
                "DEST",

                "AWB_NO",
                "AWB_DATE",

                "FLIGHT_NO",
                "AIRLINE",

                "PKGS",
                "CHR_WT",
                "RATE",

                "BASIC_FRT",
                "hawb_chr"
            ]

            preferred_order2 = [

                "TOTAL_FRT",
                "ADDON_CHR",
                "ADDON_PER_KG",

                "CPKG",
                "SLAB",
                "LODGE_MODE"
            ]

            final_columns = (

                [
                    c
                    for c in preferred_order1
                    if c in combined_df.columns
                ]

                +

                [
                    c
                    for c in combined_df.columns
                    if c not in preferred_order1 + preferred_order2
                ]

                +

                [
                    c
                    for c in preferred_order2
                    if c in combined_df.columns
                ]
            )

            combined_df = combined_df[
                final_columns
            ]

            # STORE SESSION

            st.session_state["bhagwati_df"] = combined_df

        # =====================================================
        # LOAD SESSION DATA
        # =====================================================

        combined_df = st.session_state["bhagwati_df"]

        st.success(
            "Processing Complete ✅"
        )

        # =====================================================
        # DASHBOARD
        # =====================================================

        filtered_df = show_dashboard(
            combined_df
        )

        st.dataframe(
            filtered_df,
            use_container_width=True
        )

        # =====================================================
        # EXPORT
        # =====================================================

        output = BytesIO()

        combined_df.to_excel(
            output,
            index=False
        )

        output.seek(0)

        wb = load_workbook(output)

        ws = wb.active

        # HEADER STYLE

        header_fill = PatternFill(
            start_color="1F4E79",
            end_color="1F4E79",
            fill_type="solid"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        for cell in ws[1]:

            cell.fill = header_fill
            cell.font = header_font

        # RED HIGHLIGHT

        red_fill = PatternFill(
            start_color="FFD6D6",
            end_color="FFD6D6",
            fill_type="solid"
        )

        column_colors = {
            "CHR_WT": red_fill,
            "TOTAL_FRT": red_fill,
            "CPKG": red_fill
        }

        col_index = {
            cell.value: idx + 1
            for idx, cell in enumerate(ws[1])
        }

        for col_name, fill in column_colors.items():

            if col_name in col_index:

                col_letter = get_column_letter(
                    col_index[col_name]
                )

                for row in range(
                    2,
                    ws.max_row + 1
                ):

                    ws[f"{col_letter}{row}"].fill = fill

        final_output = BytesIO()

        wb.save(final_output)

        final_output.seek(0)

        st.download_button(
            "Download Excel",
            data=final_output,
            file_name="bhagwati_processed.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )