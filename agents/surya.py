import streamlit as st
import pandas as pd
import pdfplumber
 
from io import BytesIO
from openpyxl.styles import PatternFill, Font
 
from utils.dashboard import show_dashboard
 
 
def run():
 
    st.title("📦 SURYA Processor")
 
    # =====================================================
    # FILE UPLOAD
    # =====================================================
 
    uploaded_files = st.file_uploader(
        "Upload Surya PDFs",
        type=["pdf"],
        accept_multiple_files=True,
        key="surya_upload"
    )
 
    # =====================================================
    # EXTRACTOR
    # =====================================================
 
    def extract_awb(file_bytes, section_name):
 
        rows = []
 
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
 
            text = ""
 
            for page in pdf.pages:
 
                t = page.extract_text()
 
                if t:
                    text += t + "\n"
 
            lines = text.split("\n")
 
            start = False
 
            for line in lines:
 
                if section_name in line:
 
                    start = True
                    continue
 
                if start:
 
                    if line.strip() == "":
                        continue
 
                    if "Total" in line:
                        continue
 
                    parts = line.split()
 
                    try:
 
                        if len(parts) >= 12:
 
                            rows.append({
 
                                "Awb No": parts[1],
                                "Awb Date": parts[2],
                                "Dest": parts[3],
                                "Flight": parts[4],
                                "Pkts": parts[5],
                                "Charge Wt.": parts[6],
                                "Rate": parts[7],
                                "Basic Freight": parts[9],
                                "Other Charges": parts[10],
                                "Total": parts[11],
                                "Origin": "DEL"
                            })
 
                    except:
                        continue
 
        return pd.DataFrame(rows)
 
    # =====================================================
    # HELPERS
    # =====================================================
 
    def clean_awb(x):
 
        if pd.isna(x):
            return None
 
        return "".join(
            filter(
                str.isdigit,
                str(x)
            )
        )
 
    def get_airline(f):
 
        if pd.isna(f):
            return None
 
        code = str(f)[:2].upper()
 
        return {
            "AI": "AirIndia",
            "IX": "AirIndia",
            "6E": "Indigo",
            "SG": "SpiceJet",
            "QP": "Akasa"
        }.get(code, "Other")
 
    def get_slab(wt):
 
        if pd.isna(wt):
            return None
 
        wt = float(wt)
 
        if wt < 45:
            return "Less than 45Kg"
 
        elif wt < 100:
            return "45Kg +"
 
        elif wt < 250:
            return "100Kg +"
 
        elif wt < 300:
            return "250Kg +"
 
        elif wt < 500:
            return "300Kg +"
 
        elif wt < 1000:
            return "500Kg +"
 
        return "1000Kg +"
 
    # =====================================================
    # CALCULATIONS
    # =====================================================
 
    def add_calculations(df):
 
        numeric_cols = [
            "Charge Wt.",
            "Total",
            "Basic Freight",
            "Rate",
            "Pkts",
            "Other Charges"
        ]
 
        for col in numeric_cols:
 
            df[col] = pd.to_numeric(
                df[col],
                errors="coerce"
            )
 
        # BASIC FREIGHT
 
        df["Basic Freight"] = df.apply(
            lambda r:
            r["Rate"] * r["Charge Wt."]
            if pd.isna(r["Basic Freight"])
            else r["Basic Freight"],
            axis=1
        )
 
        # OTHER CHARGES
 
        df["Other Charges"] = (
            df["Other Charges"]
            .fillna(0)
        )
 
        # TOTAL
 
        df["Total"] = df.apply(
            lambda r:
            r["Basic Freight"]
            + r["Other Charges"]
            if pd.isna(r["Total"])
            else r["Total"],
            axis=1
        )
 
        # ADDONS
 
        df["AddOns"] = (
            df["Total"]
            - df["Basic Freight"]
        )
 
        df["AddOn/Kg"] = (
            df["AddOns"]
            / df["Charge Wt."]
        ).round(2)
 
        df["CPKG"] = (
            df["Total"]
            / df["Charge Wt."]
        ).round(2)
 
        return df
 
    # =====================================================
    # FINAL FORMAT
    # =====================================================
 
    def build_index_format(df, file_name):
 
        df["Awb Date"] = pd.to_datetime(
            df["Awb Date"]
            .astype(str)
            .str.strip(),
            errors="coerce",
            dayfirst=True
        )
 
        df["AWB_DATE"] = (
            df["Awb Date"]
            .dt.strftime("%d-%b-%Y")
        )
 
        df["AIRLINE"] = (
            df["Flight"]
            .apply(get_airline)
        )
 
        df["SLAB"] = (
            df["Charge Wt."]
            .apply(get_slab)
        )
 
        df["OD_PAIR"] = (
            df["Origin"]
            + "-"
            + df["Dest"]
        )
 
        df["INVOICE_NO"] = (
            file_name.replace(".pdf", "")
        )
 
        df["AGENT"] = "SURYA"
 
        df["TRNSPT_MODE"] = "AIR"
 
        df["AWB_NO"] = df["Awb No"]
 
        df["FLIGHT_NO"] = df["Flight"]
 
        df["PKGS"] = df["Pkts"]
 
        df["CHR_WT"] = df["Charge Wt."]
 
        df["RATE"] = df["Rate"]
 
        df["BASIC_FRT"] = df["Basic Freight"]
 
        df["OTHER_CHR"] = df["Other Charges"]
 
        df["TOTAL_FRT"] = df["Total"]
 
        df["ADDON_CHR"] = df["AddOns"]
 
        df["ADDON_PER_KG"] = df["AddOn/Kg"]
 
        df["LODGE_MODE"] = df["AWB_NO"].apply(
            lambda x:
            "Direct"
            if len(str(x)) == 11
            else "Console"
        )
 
        df["ORIGIN"] = df["Origin"]
 
        df["DEST"] = df["Dest"]
 
        df["BILL_PERIOD"] = (
            df["Awb Date"]
            .dt.day
            .apply(
                lambda x:
                "FFN"
                if pd.notna(x)
                and x <= 15
                else "SFN"
            )
        )
 
        final_cols = [
 
            "INVOICE_NO",
            "BILL_PERIOD",
            "AGENT",
            "TRNSPT_MODE",
 
            "OD_PAIR",
            "ORIGIN",
            "DEST",
 
            "AWB_NO",
            "AWB_DATE",
 
            "FLIGHT_NO",
            "AIRLINE",
 
            "PKGS",
            "CHR_WT",
 
            "RATE",
            "BASIC_FRT",
 
            "OTHER_CHR",
            "TOTAL_FRT",
 
            "CPKG",
 
            "ADDON_CHR",
            "ADDON_PER_KG",
 
            "SLAB",
            "LODGE_MODE"
        ]
 
        for col in final_cols:
 
            if col not in df.columns:
                df[col] = None
 
        return df[final_cols]
 
    # =====================================================
    # MAIN
    # =====================================================
 
    if uploaded_files:
 
        # =====================================================
        # CLEAR BUTTON
        # =====================================================
 
        #if st.button("Clear Processed Data"):
 
            #if "surya_df" in st.session_state:
                #del st.session_state["surya_df"]
 
            #st.rerun()
 
        # =====================================================
        # PROCESS ONLY ONCE
        # =====================================================
 
        if "surya_df" not in st.session_state:
 
            final_df = pd.DataFrame()
 
            for file in uploaded_files:
 
                st.write(
                    f"Processing: {file.name}"
                )
 
                file_bytes = file.read()
 
                df = pd.concat([
 
                    extract_awb(
                        file_bytes,
                        "MAWB"
                    ),
 
                    extract_awb(
                        file_bytes,
                        "HAWB"
                    )
 
                ], ignore_index=True)
 
                if df.empty:
                    continue
 
                # CLEAN AWB
 
                df["Awb No"] = (
                    df["Awb No"]
                    .apply(clean_awb)
                )
 
                df = df[
                    df["Awb No"].notna()
                ]
 
                # CALCULATIONS
 
                df = add_calculations(df)
 
                # FINAL FORMAT
 
                df = build_index_format(
                    df,
                    file.name
                )
 
                final_df = pd.concat(
                    [final_df, df],
                    ignore_index=True
                )
 
            if final_df.empty:
 
                st.warning(
                    "No valid data found"
                )
 
                return
 
            final_df = (
                final_df
                .drop_duplicates(
                    subset=["AWB_NO"]
                )
            )
 
            # STANDARD DASHBOARD COLUMNS
 
            if "TOTAL_FRT" not in final_df.columns:
                final_df["TOTAL_FRT"] = 0
 
            if "ORIGIN" not in final_df.columns:
                final_df["ORIGIN"] = "DEL"
 
            if "DEST" not in final_df.columns:
                final_df["DEST"] = "UNKNOWN"
 
            if "CHR_WT" not in final_df.columns:
                final_df["CHR_WT"] = 0
 
            # STORE SESSION
 
            st.session_state["surya_df"] = final_df
 
        # =====================================================
        # LOAD SESSION DATA
        # =====================================================
 
        final_df = st.session_state["surya_df"]
 
        st.success(
            f"✅ {len(final_df)} rows combined"
        )
 
        # =====================================================
        # DASHBOARD
        # =====================================================
 
        filtered_df = show_dashboard(
            final_df
        )
 
        st.dataframe(
            filtered_df,
            use_container_width=True
        )
 
        # =====================================================
        # EXPORT
        # =====================================================
 
        output = BytesIO()
 
        highlight_cols = [
            "CHR_WT",
            "TOTAL_FRT",
            "CPKG"
        ]
 
        fill = PatternFill(
            start_color="FFD6D6",
            end_color="FFD6D6",
            fill_type="solid"
        )
 
        header_fill = PatternFill(
            start_color="4F81BD",
            end_color="4F81BD",
            fill_type="solid"
        )
 
        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:
 
            final_df.to_excel(
                writer,
                index=False,
                sheet_name="Combined"
            )
 
            ws = writer.sheets["Combined"]
 
            # HEADER STYLE
 
            for cell in ws[1]:
 
                cell.font = Font(
                    bold=True,
                    color="FFFFFF"
                )
 
                cell.fill = header_fill
 
            # FREEZE
 
            ws.freeze_panes = "A2"
 
            # FILTER
 
            ws.auto_filter.ref = ws.dimensions
 
            # COLUMN INDEX
 
            col_idx = {
                col: i + 1
                for i, col in enumerate(
                    final_df.columns
                )
            }
 
            # HIGHLIGHT
 
            for col in highlight_cols:
 
                if col in col_idx:
 
                    c = col_idx[col]
 
                    for r in range(
                        2,
                        len(final_df) + 2
                    ):
 
                        ws.cell(
                            row=r,
                            column=c
                        ).fill = fill
 
            # AUTO WIDTH
 
            for column_cells in ws.columns:
 
                length = max(
                    len(str(cell.value))
                    if cell.value
                    else 0
                    for cell in column_cells
                )
 
                ws.column_dimensions[
                    column_cells[0].column_letter
                ].width = length + 3
 
        output.seek(0)
 
        st.download_button(
            "📥 Download Excel",
            data=output,
            file_name="SURYA_INDEX_OUTPUT.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
 
 
  