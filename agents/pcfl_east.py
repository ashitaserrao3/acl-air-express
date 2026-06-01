import streamlit as st
import pandas as pd
import openpyxl
import re
 
from io import BytesIO
from openpyxl.styles import PatternFill, Font
from utils.dashboard import show_dashboard
 
 
def run():
        
            st.title("📦 PCFL East Processor")
        
            uploaded_files = st.file_uploader(
                "Upload PCFL East Excel files",
                type=["xlsx"],
                accept_multiple_files=True,
                key="pcfl_east"
            )
        
            # =================================================
            # CONFIG
            # =================================================
            AIRLINE_MAP = {
                "AI": "AirIndia",
                "IX": "AirIndia",
                "6E": "Indigo",
                "SG": "Spicejet",
                "S5": "StarAir",
                "QP": "Akasa"
            }
        
            # =================================================
            # HELPERS
            # =================================================
            def find_col(df, keyword):
        
                keyword = re.sub(
                    r"[^A-Z0-9]",
                    "",
                    keyword.upper()
                )
        
                for c in df.columns:
        
                    col = re.sub(
                        r"[^A-Z0-9]",
                        "",
                        str(c).upper()
                    )
        
                    if keyword == col:
                        return c
        
                for c in df.columns:
        
                    col = re.sub(
                        r"[^A-Z0-9]",
                        "",
                        str(c).upper()
                    )
        
                    if keyword in col:
                        return c
        
                return None
        
            def clean_awb(val):
        
                if pd.isna(val):
                    return None
        
                cleaned = re.sub(
                    r"\D",
                    "",
                    str(val)
                )
        
                return cleaned if cleaned else None
        
            def clean_weight(val):
        
                try:
        
                    if pd.isna(val):
                        return None
        
                    if str(val).strip().upper() == "M":
                        return "M"
        
                    return float(val)
        
                except:
                    return None
        
            def get_airline(flight):
        
                if pd.isna(flight):
                    return None
        
                f = (
                    str(flight)
                    .upper()
                    .replace(" ", "")
                )
        
                f = re.sub(
                    r"[^A-Z0-9]",
                    "",
                    f
                )
        
                match = re.match(
                    r"([A-Z0-9]{2})",
                    f
                )
        
                code = match.group(1) if match else ""
        
                return AIRLINE_MAP.get(code, "Other")
        
            def extract_origin(invoice):
        
                if pd.isna(invoice):
                    return None
        
                parts = str(invoice).split("-")
        
                if len(parts) >= 2:
        
                    origin = (
                        parts[0]
                        .strip()
                        .upper()
                    )
        
                    if origin.isalpha():
                        return origin
        
                return None
        
            def get_slab(wt):
        
                try:
        
                    if pd.isna(wt):
                        return None
        
                    wt = float(wt)
        
                    if wt < 45:
                        return "Less than 45Kg"
        
                    elif wt < 100:
                        return "45Kg +"
        
                    elif wt < 250:
                        return "100Kg +"
        
                    elif wt < 300:
                        return "250Kg +"
        
                    elif wt < 500:
                        return "300Kg +"
        
                    elif wt < 1000:
                        return "500Kg +"
        
                    return "1000Kg +"
        
                except:
                    return None
        
            def compute_mode(row):
        
                def val(col):
        
                    v = row.get(col, 0)
        
                    if pd.isna(v):
                        return 0
        
                    try:
                        return float(v)
        
                    except:
                        return 0
        
                awbdo = val("awbdo")
                fsc = val("fsc")
                ocdc = val("ocdc")
                tsp = val("tsp")
        
                if (
                    abs(awbdo - 50) < 0.01
                    and
                    (fsc + ocdc + tsp) == 0
                ):
                    return "Console"
        
                return "Direct"
        
            def get_bill_period(date):
        
                if pd.isna(date):
                    return None
        
                try:
        
                    day = pd.to_datetime(date).day
        
                    return (
                        "FFN"
                        if 1 <= day <= 15
                        else "SFN"
                    )
        
                except:
                    return None
        
            # =================================================
            # SAFE EXCEL LOADER
            # =================================================
            def load_excel_with_header(file):
        
                try:
        
                    if file.size == 0:
        
                        st.warning(
                            f"{file.name} is empty"
                        )
        
                        return None
        
                    xls = pd.ExcelFile(file)
        
                    if not xls.sheet_names:
        
                        st.warning(
                            f"No sheets found in {file.name}"
                        )
        
                        return None
        
                    sheet = xls.sheet_names[0]
        
                    df_raw = pd.read_excel(
                        xls,
                        sheet_name=sheet,
                        header=None
                    )
        
                    header_row = None
        
                    for i, row in df_raw.iterrows():
        
                        text = " ".join(
                            str(x).upper()
                            for x in row
                            if pd.notna(x)
                        )
        
                        required_headers = [
                            "AWB",
                            "DATE"
                        ]
        
                        if all(
                            h in text
                            for h in required_headers
                        ):
        
                            header_row = i
                            break
        
                    if header_row is None:
        
                        st.warning(
                            f"Header row not found in {file.name}"
                        )
        
                        return None
        
                    return pd.read_excel(
                        xls,
                        sheet_name=sheet,
                        header=header_row
                    )
        
                except Exception as e:
        
                    st.error(
                        f"Error reading {file.name}: {str(e)}"
                    )
        
                    return None
        
            # =================================================
            # PROCESS FILES
            # =================================================
            if uploaded_files:
        
                all_dfs = []
        
                progress = st.progress(0)
        
                with st.spinner(
                    "Processing Excel files..."
                ):
        
                    for i, file in enumerate(uploaded_files):
        
                        try:
        
                            st.write(
                                f"📄 Processing: {file.name}"
                            )
        
                            df = load_excel_with_header(file)
        
                            if df is None:
                                continue
        
                            awb_col = find_col(df, "AWB")
        
                            if not awb_col:
        
                                st.warning(
                                    f"AWB column not found in {file.name}"
                                )
        
                                continue
        
                            df_clean = pd.DataFrame()
        
                            def safe_get(keyword):
        
                                col = find_col(df, keyword)
        
                                if col:
                                    return df[col]
        
                                return None
        
                            df_clean["awb_no"] = safe_get("AWB")
        
                            date_col = safe_get("DATE")
        
                            if date_col is not None:
        
                                df_clean["awb_date"] = (
                                    pd.to_datetime(
                                        date_col,
                                        errors="coerce"
                                    ).dt.date
                                )
        
                            else:
                                df_clean["awb_date"] = None
        
                            df_clean["dest"] = safe_get("DEST")
                            df_clean["flight_no"] = safe_get("FLIGHT")
                            df_clean["pkgs"] = safe_get("PKT")
                            df_clean["chr_wt"] = safe_get("WGT")
        
                            df_clean["basic_frt"] = safe_get("BASIC")
                            df_clean["rate"] = safe_get("RATE")
        
                            df_clean["total_frt"] = safe_get("GROSS")
        
                            df_clean["awbdo"] = safe_get("AWBDO")
                            df_clean["fsc"] = safe_get("FSC")
                            df_clean["ocdc"] = safe_get("OCDC")
                            df_clean["tsp"] = safe_get("TSP")
        
                            df_clean["ssp"] = safe_get("SSP")
                            df_clean["ssp_amt"] = safe_get("SSP AMT")
                            df_clean["dis%"] = safe_get("DIS")
                            df_clean["disc"] = safe_get("DISC")
        
                            df_clean["cgst"] = safe_get("CGST")
                            df_clean["sgst"] = safe_get("SGST")
                            df_clean["igst"] = safe_get("IGST")
        
                            df_clean["amount"] = safe_get("AMOUNT")
        
                            df_clean["awb_no"] = (
                                df_clean["awb_no"]
                                .apply(clean_awb)
                            )
        
                            df_clean = df_clean[
                                df_clean["awb_no"].notna()
                            ]
        
                            df_clean["chr_wt"] = (
                                df_clean["chr_wt"]
                                .apply(clean_weight)
                            )
        
                            df_clean["airline"] = (
                                df_clean["flight_no"]
                                .apply(get_airline)
                            )
        
                            df_clean["slab"] = (
                                df_clean["chr_wt"]
                                .apply(get_slab)
                            )
        
                            df_clean["invoice_no"] = (
                                file.name.replace(".xlsx", "")
                            )
        
                            df_clean["origin"] = (
                                df_clean["invoice_no"]
                                .apply(extract_origin)
                            )
        
                            gross = pd.to_numeric(
                                df_clean["total_frt"],
                                errors="coerce"
                            )
        
                            wt = pd.to_numeric(
                                df_clean["chr_wt"],
                                errors="coerce"
                            )
        
                            df_clean["cpkg"] = (
                                gross
                                .div(wt)
                                .round(2)
                            )
        
                            df_clean.loc[
                                (
                                    wt.isna()
                                )
                                |
                                (
                                    wt == 0
                                ),
                                "cpkg"
                            ] = None
        
                            df_clean["lodge_mode"] = (
                                df_clean.apply(
                                    compute_mode,
                                    axis=1
                                )
                            )
        
                            df_clean["od_pair"] = (
                                df_clean["origin"]
                                .fillna("")
                                +
                                "-"
                                +
                                df_clean["dest"]
                                .fillna("")
                            )
        
                            df_clean["agent"] = "PCFL(E)"
        
                            df_clean["transport_mode"] = "Air"
        
                            df_clean["bill_period"] = (
                                df_clean["awb_date"]
                                .apply(get_bill_period)
                            )
        
                            all_dfs.append(df_clean)
        
                        except Exception as e:
        
                            st.error(
                                f"Error processing {file.name}: {str(e)}"
                            )
        
                        progress.progress(
                            (i + 1)
                            / len(uploaded_files)
                        )
        
                if all_dfs:
        
                    combined_df = pd.concat(
                        all_dfs,
                        ignore_index=True
                    )
        
                    combined_df = combined_df.drop_duplicates(
                        subset=[
                            "awb_no",
                            "awb_date",
                            "invoice_no"
                        ]
                    )
        
                    combined_df.columns = [
                        c.upper()
                        for c in combined_df.columns
                    ]


                    st.success(
                        f"✅ Combined {len(combined_df)} rows"
                    )

                    # DASHBOARD

                    filtered_df = show_dashboard(combined_df)
                        
                    st.dataframe(
                            filtered_df,
                            use_container_width=True
                                )
        
                    output = BytesIO()
        
                    with pd.ExcelWriter(
                        output,
                        engine="openpyxl"
                    ) as writer:
        
                        combined_df.to_excel(
                            writer,
                            index=False,
                            sheet_name="Combined"
                        )
        
                        ws = writer.sheets["Combined"]
        
                        ws.freeze_panes = "A2"
        
                        ws.auto_filter.ref = ws.dimensions
        
                        light_red = PatternFill(
                            start_color="FFD6D6",
                            end_color="FFD6D6",
                            fill_type="solid"
                        )
        
                        headers = {
                            cell.value: idx + 1
                            for idx, cell in enumerate(ws[1])
                        }
        
                        def highlight(col_name):
        
                            if col_name in headers:
        
                                col_idx = headers[col_name]
        
                                ws.cell(
                                    row=1,
                                    column=col_idx
                                ).fill = light_red
        
                                ws.cell(
                                    row=1,
                                    column=col_idx
                                ).font = Font(
                                    bold=True
                                )
        
                                for r in range(
                                    2,
                                    len(combined_df) + 2
                                ):
        
                                    ws.cell(
                                        row=r,
                                        column=col_idx
                                    ).fill = light_red
        
                        highlight("CHR_WT")
                        highlight("CPKG")
                        highlight("TOTAL_FRT")
        
                        for column_cells in ws.columns:
        
                            values = [
                                str(cell.value)
                                for cell in column_cells
                                if cell.value is not None
                            ]
        
                            if not values:
                                continue
        
                            length = max(
                                len(v)
                                for v in values
                            )
        
                            ws.column_dimensions[
                                column_cells[0].column_letter
                            ].width = min(
                                length + 2,
                                40
                            )
        
                    output.seek(0)
        
                    st.download_button(
                        "📥 Download Excel",
                        data=output,
                        file_name="PCFL_East_Combined.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
        
                else:
        
                    st.warning("No valid data found.")
        