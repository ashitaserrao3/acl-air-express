import streamlit as st
import pandas as pd
import re
import openpyxl
 
from io import BytesIO
 
from utils.exporter import export_excel
from utils.dashboard import show_dashboard
from utils.slab import get_slab
 
 
# ---------------- SHEET SKIP ----------------
def should_skip(sheet_name):
 
    name = " ".join(
        sheet_name.lower().split()
    )
 
    skip_patterns = [
        "bill summary",
        "way bill report",
        "waybill report"
    ]
 
    return any(
        p in name
        for p in skip_patterns
    )
 
 
# ---------------- PROCESS FILE ----------------
def process_pobc(file):
 
    wb = openpyxl.load_workbook(
        file,
        data_only=True
    )
 
    combined_df = pd.DataFrame()
 
    for sheet_name in wb.sheetnames:
 
        if should_skip(sheet_name):
            continue
 
        try:
 
            sheet = wb[sheet_name]
 
            # ---------------- UNMERGE ----------------
            for merged_range in list(
                sheet.merged_cells.ranges
            ):
 
                min_row = merged_range.min_row
                max_row = merged_range.max_row
 
                min_col = merged_range.min_col
                max_col = merged_range.max_col
 
                value = sheet.cell(
                    row=min_row,
                    column=min_col
                ).value
 
                sheet.unmerge_cells(
                    str(merged_range)
                )
 
                for r in range(
                    min_row,
                    max_row + 1
                ):
 
                    for c in range(
                        min_col,
                        max_col + 1
                    ):
 
                        sheet.cell(
                            row=r,
                            column=c
                        ).value = value
 
            # ---------------- RAW DF ----------------
            df_raw = pd.DataFrame(
                list(sheet.values)
            ).ffill()
 
            header_row = None
 
            for i, row in df_raw.iterrows():
 
                row_text = " ".join(
                    [str(x).lower() for x in row]
                )
 
                if (
                    "waybill" in row_text
                    and "route" in row_text
                    and "dest" in row_text
                ):
 
                    header_row = i
                    break
 
            if header_row is None:
                continue
 
            # ---------------- CLEAN HEADERS ----------------
            headers = (
                df_raw.iloc[header_row]
                .astype(str)
                .str.strip()
            )
 
            seen = {}
            clean_headers = []
 
            for h in headers:
 
                h = (
                    "col"
                    if h.lower() == "nan"
                    else h
                )
 
                if h in seen:
 
                    seen[h] += 1
 
                    clean_headers.append(
                        f"{h}_{seen[h]}"
                    )
 
                else:
 
                    seen[h] = 0
                    clean_headers.append(h)
 
            # ---------------- DATAFRAME ----------------
            df = pd.DataFrame(
                df_raw.values[
                    header_row + 1:
                ],
                columns=clean_headers
            )
 
            # ---------------- CLEAN COLUMNS ----------------
            df.columns = (
                df.columns.astype(str)
                .str.strip()
                .str.replace("\n", " ")
                .str.replace(
                    r"\s+",
                    " ",
                    regex=True
                )
                .str.lower()
            )
 
            # ---------------- COLUMN MAPPING ----------------
            mapping = {}
 
            for col in df.columns:
 
                c = col.lower()
 
                if "waybill" in c:
                    mapping[col] = "AWB_NO"
 
                elif "awbc" in c:
                    mapping[col] = "awbc"
 
                elif c == "date":
                    mapping[col] = "AWB_DATE"
 
                elif "route" in c:
                    mapping[col] = "FLIGHT_NO"
 
                elif c == "dest":
                    mapping[col] = "DEST"
 
                elif "srv" in c:
                    mapping[col] = "TRNSPT_MODE"
 
                elif "pcs" in c:
                    mapping[col] = "PKGS"
 
                elif "bag" in c:
                    mapping[col] = "BAGS"
 
                elif c == "wt":
                    mapping[col] = "CHR_WT"
 
                elif c == "rate":
                    mapping[col] = "RATE"
 
                elif "frt" in c:
                    mapping[col] = "BASIC_FRT"
 
                elif "revenue" in c:
                    mapping[col] = "TOTAL_FRT"
 
                elif "oth" in c:
                    mapping[col] = "oth"
 
                elif "hndc" in c:
                    mapping[col] = "hndc"
 
                elif "dunt" in c:
                    mapping[col] = "dunt"
 
                elif "tspi" in c:
                    mapping[col] = "tspi"
 
                elif "tspo" in c:
                    mapping[col] = "tspo"
 
                elif "unit" in c:
                    mapping[col] = "unit"
 
                elif "xray" in c:
                    mapping[col] = "xray"
 
                elif "srch" in c:
                    mapping[col] = "srch"
 
                else:
                    mapping[col] = c
 
            df.rename(
                columns=mapping,
                inplace=True
            )
 
            df = df.applymap(
                lambda x:
                x.strip()
                if isinstance(x, str)
                else x
            )
 
            # ---------------- DEST CLEAN ----------------
            if "DEST" in df.columns:
 
                df["DEST"] = (
                    df["DEST"]
                    .astype(str)
                    .str.replace(
                        r"^(PAF-|POBC-)",
                        "",
                        regex=True
                    )
                    .str.strip()
                )
 
            # ---------------- ORIGIN ----------------
            df["ORIGIN"] = (
                sheet_name[2:5]
                if len(sheet_name) >= 5
                else sheet_name
            )
 
            df["ORIGIN"] = df["ORIGIN"].replace(
                "XM1",
                "IXM"
            )
 
            # ---------------- AWB CLEAN ----------------
            if "AWB_NO" in df.columns:
 
                df["AWB_NO"] = (
                    df["AWB_NO"]
                    .astype(str)
                    .str.strip()
                )
 
                df["AWB_NO"] = (
                    df["AWB_NO"]
                    .str.replace(
                        r"\b(AI|6E|S5|SG|QP|IX)\b[- ]*",
                        "",
                        regex=True
                    )
                )
 
                df = df[
                    df["AWB_NO"].notna()
                    &
                    (df["AWB_NO"] != "")
                    &
                    (
                        df["AWB_NO"]
                        .str.lower() != "nan"
                    )
                    &
                    (
                        df["AWB_NO"]
                        .str.contains(r"\d", na=False)
                    )
                ]
 
            # ---------------- REMOVE BLANK ROWS ----------------
            df = df.dropna(how="all")
 
            # ---------------- NUMERIC ----------------
            num_cols = [
                "awbc",
                "oth",
                "hndc",
                "dunt",
                "tspi",
                "tspo",
                "unit",
                "xray",
                "srch",
                "CHR_WT",
                "BASIC_FRT",
                "TOTAL_FRT"
            ]
 
            for col in num_cols:
 
                if col in df.columns:
 
                    df[col] = pd.to_numeric(
                        df[col],
                        errors="coerce"
                    ).fillna(0)
 
            # ---------------- DATE CLEAN ----------------
            if "AWB_DATE" in df.columns:
 
                df["AWB_DATE"] = pd.to_datetime(
                    df["AWB_DATE"],
                    errors="coerce"
                )
 
                df = df[
                    df["AWB_DATE"].notna()
                    &
                    (
                        df["AWB_DATE"]
                        .dt.year != 1970
                    )
                ]
 
                df["AWB_DATE"] = (
                    df["AWB_DATE"]
                    .dt.date
                )
 
            else:
 
                df["AWB_DATE"] = None
 
            # ---------------- TRANSPORT MODE ----------------
            if "TRNSPT_MODE" in df.columns:
 
                df["TRNSPT_MODE"] = (
                    df["TRNSPT_MODE"]
                    .astype(str)
                    .str.strip()
                    .str.upper()
                    .replace({
                        "ACG": "AIR",
                        "RDS": "ROAD"
                    })
                )
 
            else:
 
                df["TRNSPT_MODE"] = "N/A"
 
            # ---------------- CALCULATIONS ----------------
            df["ADDON"] = (
                df["TOTAL_FRT"]
                - df["BASIC_FRT"]
            )
 
            wt = df["CHR_WT"].replace(
                0,
                pd.NA
            )
 
            df["ADDON_PER_KG"] = (
                df["ADDON"] / wt
            ).fillna(0)
 
            df["CPKG"] = (
                df["TOTAL_FRT"] / wt
            ).fillna(0)
 
            # ---------------- SLAB ----------------
            df["SLAB"] = (
                df["CHR_WT"]
                .apply(get_slab)
            )
 
            # ---------------- INVOICE ----------------
            df["INVOICE_NO"] = sheet_name
 
            # ---------------- COMBINE ----------------
            combined_df = pd.concat(
                [combined_df, df],
                ignore_index=True
            )
 
        except Exception as e:
 
            st.error(
                f"Error in sheet '{sheet_name}': {e}"
            )
 
    combined_df = combined_df.drop(
        columns=["no."],
        errors="ignore"
    )
 
    return combined_df
 
 
# ---------------- LODGE MODE ----------------
def compute_mode(row):
 
    def val(col):
 
        v = row.get(col, 0)
 
        if pd.isna(v):
            return 0
 
        try:
            return float(v)
 
        except:
            return 0
 
    transport = str(
    row.get("TRNSPT_MODE", "")
    ).strip().upper()
 
    if transport != "AIR":
        return "N/A"
 
 
    awbc = val("awbc")
    dunt = val("dunt")
    oth = val("oth")
 
    cols_a = [
        "oth",
        "hndc",
        "tspi",
        "tspo",
        "dunt",
        "unit",
        "xray",
        "srch"
    ]
 
    cond_a = (
        0 < awbc <= 125
    ) and (
        sum(val(c) for c in cols_a) == 0
    )
 
    cols_b = [
        "awbc",
        "oth",
        "hndc",
        "tspi",
        "tspo",
        "unit",
        "xray",
        "srch"
    ]
 
    cond_b = (
        0 < dunt <= 125
    ) and (
        sum(val(c) for c in cols_b) == 0
    )
 
    cols_c = [
        "awbc",
        "dunt",
        "hndc",
        "tspi",
        "tspo",
        "unit",
        "xray",
        "srch"
    ]
 
    cond_c = (
        0 < oth <= 125
    ) and (
        sum(val(c) for c in cols_c) == 0
    )
 
    return (
        "Console"
        if (cond_a or cond_b or cond_c)
        else "Direct"
    )
 
 
# ---------------- MAIN ----------------
def run():
 
    st.title("📦 POBC Processor")
 
    uploaded_files = st.file_uploader(
        "Upload POBC Excel files",
        type=["xlsx"],
        accept_multiple_files=True,
        key="pobc_upload"
    )
 
    if not uploaded_files:
        return
 
 
    # ---------------- PROCESS ONLY ONCE ----------------
    if "pobc_df" not in st.session_state:
 
        final_df = pd.DataFrame()
 
        progress = st.progress(0)
 
        for i, file in enumerate(uploaded_files):
 
            df = process_pobc(file)
 
            final_df = pd.concat(
                [final_df, df],
                ignore_index=True
            )
 
            progress.progress(
                (i + 1) / len(uploaded_files)
            )
 
        if final_df.empty:
 
            st.error(
                "❌ No valid data extracted."
            )
 
            return
 
        # ---------------- REMOVE DUPLICATES ----------------
        final_df = final_df.drop_duplicates()
 
        # ---------------- LODGE MODE ----------------
        final_df["LODGE_MODE"] = final_df.apply(
            compute_mode,
            axis=1
        )
 
        final_df["AWB_NO"] = (
            final_df["AWB_NO"]
            .astype(str)
            .str.strip()
        )
 
        # ---------------- SLASH OVERRIDE ----------------
        final_df.loc[
            final_df["AWB_NO"]
            .str.contains("/", na=False),
 
            "LODGE_MODE"
 
        ] = "Console"
 
        # ---------------- FINAL OVERRIDE ----------------
        if (
            "awbc" in final_df.columns
            and
            "hndc" in final_df.columns
        ):
 
            final_df.loc[
                (
                    final_df["ORIGIN"]
                    .astype(str)
                    .str.strip()
                    .str.upper() == "RPR"
                )
                &
                (
                    final_df["DEST"]
                    .astype(str)
                    .str.strip()
                    .str.upper() == "DEL"
                )
                &
                (
                    final_df["awbc"] == 50
                )
                &
                (
                    final_df["hndc"]
                    .fillna(0) != 0
                ),
 
                "LODGE_MODE"
 
            ] = "Console"
 
        # ---------------- STATIC ----------------
        final_df["AGENT"] = "POBC"
 
        final_df["BILL_PERIOD"] = final_df[
            "AWB_DATE"
        ].apply(
            lambda d:
            "FFN"
            if pd.notna(d)
            and d.day <= 15
            else "SFN"
        )
 
        # ---------------- COLUMN ORDER ----------------
        preferred_order = [
 
            "INVOICE_NO",
            "BILL_PERIOD",
            "AGENT",
            "TRNSPT_MODE",
 
            "OD_PAIR",
            "ORIGIN",
            "DEST",
 
            "AWB_NO",
            "AWB_DATE",
 
            "FLIGHT_NO",
 
            "PKGS",
            "CHR_WT",
 
            "BAGS",
            "RATE",
 
            "BASIC_FRT",
 
            "awbc",
            "oth",
            "hndc",
            "dunt",
            "tspi",
            "tspo",
            "unit",
            "xray",
            "srch",
 
            "TOTAL_FRT",
 
            "CPKG",
 
            "ADDON",
            "ADDON_PER_KG",
 
            "SLAB",
            "LODGE_MODE"
        ]
 
        final_df = final_df[
            [c for c in preferred_order if c in final_df.columns]
            +
            [
                c
                for c in final_df.columns
                if c not in preferred_order
            ]
        ]
 
        # ---------------- STORE SESSION ----------------
        st.session_state["pobc_df"] = final_df
 
    # ---------------- LOAD SESSION ----------------
    final_df = st.session_state["pobc_df"]
 
    # ---------------- SUCCESS ----------------
    st.success(
        f"✅ Total rows: {len(final_df)}"
    )
 
    # ---------------- DASHBOARD ----------------
    filtered_df = show_dashboard(final_df)
 
    # ---------------- DATAFRAME ----------------
    st.dataframe(
        filtered_df,
        use_container_width=True
    )
 
    # ---------------- EXPORT ----------------
    output = export_excel(
        final_df,
        sheet_name="POBC"
    )
 
    st.download_button(
        "📥 Download Excel",
        data=output,
        file_name="POBC_Combined.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
 