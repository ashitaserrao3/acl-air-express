import streamlit as st
import pandas as pd
import pdfplumber
import re
from io import BytesIO
from openpyxl.styles import PatternFill
from utils.dashboard import show_dashboard


def run():

    st.set_page_config(
        page_title="PALM PDF Extractor",
        layout="wide"
    )

    st.title("📦 PALM Processor")

    # ---------------------------------------------------
    # SLAB LOGIC
    # ---------------------------------------------------
    def get_slab(wt):

        if pd.isna(wt):
            return None

        wt = float(wt)

        if wt < 45:
            return "Less than 45Kg"
        elif wt < 100:
            return "45Kg +"
        elif wt < 250:
            return "100Kg +"
        elif wt < 300:
            return "250Kg +"
        elif wt < 500:
            return "300Kg +"
        elif wt < 1000:
            return "500Kg +"

        return "1000Kg +"

    # ---------------------------------------------------
    # BILL PERIOD LOGIC
    # ---------------------------------------------------
    def get_bill_period(date_str):

        try:
            dt = pd.to_datetime(
                date_str,
                format="%d.%m.%y",
                errors="coerce"
            )

            if pd.isna(dt):
                return None

            return "FFN" if dt.day <= 15 else "SFN"

        except:
            return None

    # ---------------------------------------------------
    # EXTRACT DATA FUNCTION
    # ---------------------------------------------------
    def extract_data(text):

        extracted_rows = []

        # ------------------------------------------------
        # FIND INVOICE NUMBER
        # ------------------------------------------------
        invoice_no = None

        invoice_match = re.search(
            r'(C/\d+/\d{2}-\d{2})',
            text
        )

        if invoice_match:
            invoice_no = invoice_match.group(1)

        # ------------------------------------------------
        # FIND TABLE ROWS
        # ------------------------------------------------
        lines = text.split("\n")

        pattern = re.compile(
            r'^\d+\s+'
            r'(\d{1,2}.\d{1,2}.\d{2})\s+'   # DATE
            r'(\d+)\s+'                     # AWB NO
            r'([A-Z]{3})\s+'                # DEST
            r'(\d+)\s+'                     # PKGS
            r'([\d.]+)\s+'                  # WT
            r'([\d.]+)\s+'                  # RATE
            r'([\d.]+)'                     # TOTAL PRICE
        )

        for line in lines:

            match = pattern.match(line.strip())

            if match:

                awb_date = match.group(1)
                awb_no = match.group(2)
                dest = match.group(3)

                pkgs = pd.to_numeric(
                    match.group(4),
                    errors="coerce"
                )

                chr_wt = pd.to_numeric(
                    match.group(5),
                    errors="coerce"
                )

                rate = pd.to_numeric(
                    match.group(6),
                    errors="coerce"
                )

                total_frt = pd.to_numeric(
                    match.group(7),
                    errors="coerce"
                )

                # ------------------------------------------------
                # BASIC FREIGHT
                # ------------------------------------------------
                basic_frt = total_frt

                # ------------------------------------------------
                # ADDON CHARGE
                # ------------------------------------------------
                addon_chr = total_frt - basic_frt

                # ------------------------------------------------
                # CPKG
                # ------------------------------------------------
                cpkg = None

                if chr_wt and chr_wt != 0:
                    cpkg = round(
                        total_frt / chr_wt,
                        2
                    )

                # ------------------------------------------------
                # ADDON PER KG
                # ------------------------------------------------
                addon_per_kg = None

                if chr_wt and chr_wt != 0:
                    addon_per_kg = round(
                        addon_chr / chr_wt,
                        2
                    )

                # ------------------------------------------------
                # FINAL ROW
                # ------------------------------------------------
                row = {
                    "INVOICE_NO": invoice_no,
                    "BILL_PERIOD": get_bill_period(awb_date),
                    "AGENT": "PALM",
                    "TRNSPT_MODE": "AIR",
                    "OD_PAIR": f"HYD-{dest}",
                    "ORIGIN": "HYD",
                    "DEST": dest,
                    "AWB_NO": awb_no,
                    "AWB_DATE": awb_date,
                    "PKGS": pkgs,
                    "CHR_WT": chr_wt,
                    "RATE": rate,
                    "BASIC_FRT": basic_frt,
                    "TOTAL_FRT": total_frt,
                    "CPKG": cpkg,
                    "ADDON_CHR": addon_chr,
                    "ADDON_PER_KG": addon_per_kg,
                    "SLAB": get_slab(chr_wt),
                    "LODGE_MODE": "Console"
                }

                extracted_rows.append(row)

        return pd.DataFrame(extracted_rows)

    # ---------------------------------------------------
    # FILE UPLOAD
    # ---------------------------------------------------
    uploaded_files = st.file_uploader(
        "Upload PALM PDF Files",
        type=["pdf"],
        accept_multiple_files=True,
        key="palm_upload"
    )

    # ---------------------------------------------------
    # PROCESS FILES
    # ---------------------------------------------------
    if uploaded_files:

        # =====================================================
        # CLEAR BUTTON
        # =====================================================
        #if st.button("Clear Processed Data"):

            #if "palm_df" in st.session_state:
                #del st.session_state["palm_df"]

            #st.rerun()

        # =====================================================
        # PROCESS ONLY ONCE
        # =====================================================
        if "palm_df" not in st.session_state:

            combined_df = pd.DataFrame()

            progress = st.progress(0)

            for i, file in enumerate(uploaded_files):

                st.write(f"Processing: {file.name}")

                try:

                    full_text = ""

                    with pdfplumber.open(file) as pdf:

                        for page in pdf.pages:

                            text = page.extract_text()

                            if text:
                                full_text += "\n" + text

                    df = extract_data(full_text)

                    combined_df = pd.concat(
                        [combined_df, df],
                        ignore_index=True
                    )

                except Exception as e:

                    st.error(
                        f"Error reading {file.name}: {e}"
                    )

                progress.progress(
                    (i + 1) / len(uploaded_files)
                )

            # ---------------------------------------------------
            # FINAL OUTPUT
            # ---------------------------------------------------
            if combined_df.empty:

                st.warning("No data extracted.")
                return

            # STANDARD DASHBOARD COLUMNS
            if "TOTAL_FRT" not in combined_df.columns:
                combined_df["TOTAL_FRT"] = 0

            if "ORIGIN" not in combined_df.columns:
                combined_df["ORIGIN"] = "HYD"

            if "DEST" not in combined_df.columns:
                combined_df["DEST"] = "UNKNOWN"

            if "CHR_WT" not in combined_df.columns:
                combined_df["CHR_WT"] = 0

            final_columns = [
                "INVOICE_NO",
                "BILL_PERIOD",
                "AGENT",
                "TRNSPT_MODE",
                "OD_PAIR",
                "ORIGIN",
                "DEST",
                "AWB_NO",
                "AWB_DATE",
                "PKGS",
                "CHR_WT",
                "RATE",
                "BASIC_FRT",
                "TOTAL_FRT",
                "CPKG",
                "ADDON_CHR",
                "ADDON_PER_KG",
                "SLAB",
                "LODGE_MODE"
            ]

            for col in final_columns:

                if col not in combined_df.columns:
                    combined_df[col] = None

            combined_df = combined_df[final_columns]

            # STORE SESSION
            st.session_state["palm_df"] = combined_df

        # =====================================================
        # LOAD SESSION DATA
        # =====================================================
        combined_df = st.session_state["palm_df"]

        st.success(
            f"✅ Extracted {len(combined_df)} rows"
        )

        # =====================================================
        # DASHBOARD
        # =====================================================
        filtered_df = show_dashboard(combined_df)

        st.dataframe(
            filtered_df,
            use_container_width=True
        )

        # ---------------------------------------------------
        # EXPORT EXCEL
        # ---------------------------------------------------
        output = BytesIO()

        highlight_cols = [
            "CHR_WT",
            "TOTAL_FRT",
            "CPKG"
        ]

        fill = PatternFill(
            start_color="FFD6D6",
            end_color="FFD6D6",
            fill_type="solid"
        )

        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:

            combined_df.to_excel(
                writer,
                index=False,
                sheet_name="PALM_DATA"
            )

            worksheet = writer.sheets["PALM_DATA"]

            # -----------------------------------------------
            # COLUMN POSITION MAP
            # -----------------------------------------------
            col_idx = {
                col: idx + 1
                for idx, col in enumerate(combined_df.columns)
            }

            # -----------------------------------------------
            # HIGHLIGHT CELLS
            # -----------------------------------------------
            for col in highlight_cols:

                if col in col_idx:

                    col_pos = col_idx[col]

                    for row in range(
                        2,
                        len(combined_df) + 2
                    ):

                        worksheet.cell(
                            row=row,
                            column=col_pos
                        ).fill = fill

        output.seek(0)

        st.download_button(
            label="📥 Download Excel",
            data=output,
            file_name="PALM_Extracted.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )